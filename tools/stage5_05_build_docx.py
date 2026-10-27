"""Build the manual from its Markdown using the selected bundled python-docx runtime."""
from pathlib import Path
import argparse,hashlib,json,re
from docx import Document
from docx.shared import Pt,Mm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
def fonts(style,code=False):
    style.font.name='Consolas' if code else 'Times New Roman'
    style.font.size=Pt(9.5 if code else 12);style.font.color.rgb=RGBColor(0,0,0)
    prop=style.element.get_or_add_rPr();family=prop.find(qn('w:rFonts'))
    if family is None:family=OxmlElement('w:rFonts');prop.insert(0,family)
    for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:family.attrib.pop(qn('w:'+key),None)
    for key,value in [('ascii',style.font.name),('hAnsi',style.font.name),('eastAsia','楷体'),('cs','Times New Roman')]:family.set(qn('w:'+key),value)
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'en-US');lang.set(qn('w:eastAsia'),'zh-CN');prop.append(lang)

def inline(paragraph,text):
    # Formatting tokens are removed only outside source fences.
    tokens=re.split(r'(\*\*[^*]+\*\*|`[^`]+`|!?\[[^\]]+\]\([^)]+\))',text)
    for token in tokens:
        if not token:continue
        bold=token.startswith('**') and token.endswith('**')
        code=token.startswith('`') and token.endswith('`')
        link=re.fullmatch(r'!?\[([^\]]+)\]\(([^)]+)\)',token)
        value=token[2:-2] if bold else token[1:-1] if code else link[1] if link else token
        run=paragraph.add_run(value)
        if bold:run.bold=True
        if code:run.font.name='Consolas';run.font.size=Pt(10)

def zero_indent(paragraph_format,element):
    paragraph_format.first_line_indent=Pt(0)
    ind=element.get_or_add_pPr().find(qn('w:ind'))
    if ind is not None:ind.set(qn('w:firstLineChars'),'0')

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--markdown',type=Path,required=True);parser.add_argument('--output',type=Path,required=True);a=parser.parse_args()
    source=a.markdown.resolve();output=a.output.resolve()
    assert source.is_relative_to(ROOT) and output.is_relative_to(ROOT)
    text=source.read_text(encoding='utf-8');lines=text.splitlines();doc=Document()
    section=doc.sections[0];section.page_width=Mm(215.9);section.page_height=Mm(279.4)
    section.top_margin=section.bottom_margin=Mm(20);section.left_margin=section.right_margin=Mm(20)
    normal=doc.styles['Normal'];fonts(normal);normal.paragraph_format.line_spacing=1.5
    normal.paragraph_format.first_line_indent=Pt(24);normal.paragraph_format.space_after=Pt(4)
    normal.paragraph_format.widow_control=True
    # Word uses firstLineChars=200 for two full body characters.
    ind=normal.element.get_or_add_pPr().find(qn('w:ind'))
    if ind is not None:ind.set(qn('w:firstLineChars'),'200')
    for name in ['Title','Heading 1','Heading 2','Heading 3','Heading 4','Heading 5','Heading 6']:
        style=doc.styles[name];fonts(style);style.font.bold=True
        style.paragraph_format.first_line_indent=Pt(0);style.paragraph_format.line_spacing=1.5
        style.paragraph_format.space_before=Pt(12);style.paragraph_format.space_after=Pt(6)
        style.paragraph_format.keep_with_next=True
        zero_indent(style.paragraph_format,style.element)
    code=doc.styles.add_style('Source Code',1);fonts(code,True)
    code.paragraph_format.line_spacing=1.05;code.paragraph_format.space_before=Pt(0);code.paragraph_format.space_after=Pt(0)
    code.paragraph_format.first_line_indent=Pt(0);code.paragraph_format.keep_together=False;code.paragraph_format.widow_control=False
    zero_indent(code.paragraph_format,code.element)
    for name in ['List Bullet','List Number']:
        style=doc.styles[name];fonts(style);style.paragraph_format.line_spacing=1.5;style.paragraph_format.first_line_indent=Pt(0)
        zero_indent(style.paragraph_format,style.element)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.paragraph_format.first_line_indent=Pt(0)
    zero_indent(footer.paragraph_format,footer._p)
    page=OxmlElement('w:fldSimple');page.set(qn('w:instr'),'PAGE');footer._p.append(page)
    doc.core_properties.author='';doc.core_properties.title='RK3576 端侧多模态系统集成与性能优化工程手册'
    doc.core_properties.subject='Stage5 5.1 to 5.5 verified engineering records'
    index=0;code_blocks=[];image_count=0
    while index<len(lines):
        line=lines[index]
        if line.startswith('```'):
            block=[];index+=1
            while index<len(lines) and not lines[index].startswith('```'):
                block.append(lines[index]);doc.add_paragraph(lines[index],style='Source Code');index+=1
            assert index<len(lines),'unclosed source fence';code_blocks.append(block);doc.add_paragraph()
        elif line=='<!-- WORD_TOC -->':
            paragraph=doc.add_paragraph();zero_indent(paragraph.paragraph_format,paragraph._p)
            field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u');paragraph._p.append(field)
        elif line=='<!-- PAGE_BREAK -->':doc.add_page_break()
        elif line.startswith('<!--'):pass
        elif re.match(r'^#{1,6} ',line):
            marker,title=line.split(' ',1);paragraph=doc.add_paragraph(style='Title' if len(marker)==1 else 'Heading '+str(len(marker)-1))
            inline(paragraph,title)
        elif line.startswith('!['):
            match=re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)',line)
            if match:
                path=(source.parent/match[2]).resolve();assert path.is_file(),path
                paragraph=doc.add_paragraph();paragraph.paragraph_format.first_line_indent=Pt(0);paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
                zero_indent(paragraph.paragraph_format,paragraph._p)
                paragraph.add_run().add_picture(str(path),width=Mm(164));image_count+=1
                caption=doc.add_paragraph(match[1]);caption.paragraph_format.first_line_indent=Pt(0);caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
                zero_indent(caption.paragraph_format,caption._p)
            else:inline(doc.add_paragraph(),line)
        elif line.startswith('|'):
            rows=[]
            while index<len(lines) and lines[index].startswith('|'):
                row=re.split(r'(?<!\\)\|',lines[index].strip().strip('|'))
                if not all(re.fullmatch(r'\s*:?-+:?\s*',cell) for cell in row):rows.append([x.strip().replace('\\|','|') for x in row])
                index+=1
            columns=max(map(len,rows));table=doc.add_table(rows=0,cols=columns);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
            lengths=[max(4,min(28,max((len(row[k]) if k<len(row) else 0) for row in rows))) for k in range(columns)]
            widths=[170*n/sum(lengths) for n in lengths]
            if rows[0][:2]==['模块','真实模型或Runtime']:
                widths=[22]+[148*n/sum(lengths[1:]) for n in lengths[1:]]
            elif columns==6 and rows[0][0]=='case' and rows[0][1]=='信号':
                widths=[50,24,27,20,27,22]
            for k,column in enumerate(table.columns):column.width=Mm(widths[k])
            borders=OxmlElement('w:tblBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                node=OxmlElement('w:'+edge);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9');borders.append(node)
            table._tbl.tblPr.append(borders)
            for n,row in enumerate(rows):
                cells=table.add_row().cells
                if n==0:
                    repeated=OxmlElement('w:tblHeader');table.rows[-1]._tr.get_or_add_trPr().append(repeated)
                for k,cell in enumerate(cells):
                    cell.width=Mm(widths[k]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
                    zero_indent(p.paragraph_format,p._p)
                    inline(p,row[k] if k<len(row) else '')
                    if n==0 or re.fullmatch(r'[-+0-9.%/ ]+',row[k] if k<len(row) else ''):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    if n==0:
                        for run in p.runs:run.bold=True
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EDF2');cell._tc.get_or_add_tcPr().append(shade)
            doc.add_paragraph();continue
        elif line.strip():
            match=re.match(r'^(- |\d+\. )(.*)',line)
            paragraph=doc.add_paragraph(style='List Bullet' if match and match[1]=='- ' else 'List Number' if match else 'Normal')
            inline(paragraph,match[2] if match else line.removeprefix('> '))
        index+=1
    # The default Word template can contain a blue Title border. Remove it,
    # along with theme color overrides, before native Word reopens the artifact.
    for element in [doc.styles.element,doc.element]:
        for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for style in doc.styles:
        if style.name=='Title' or style.name.startswith('Heading'):
            prop=style.element.get_or_add_rPr();color=prop.find(qn('w:color'))
            if color is not None:
                color.set(qn('w:val'),'000000')
                for key in ['themeColor','themeTint','themeShade']:color.attrib.pop(qn('w:'+key),None)
    settings=doc.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
    reopened=Document(output);code_lines=[p.text for p in reopened.paragraphs if p.style.name=='Source Code']
    expected=[line for block in code_blocks for line in block];assert code_lines==expected,'source code text was altered'
    result={'markdown_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'docx_sha256':hashlib.sha256(output.read_bytes()).hexdigest(),'source_code_blocks':len(code_blocks),'source_code_lines':len(expected),'all_source_code_lines_preserved':True,'images':image_count,'paragraphs':len(reopened.paragraphs),'tables':len(reopened.tables),'fonts':'body Chinese KaiTi and Latin Times New Roman 12pt; code Consolas 9.5pt','body_spacing':1.5,'body_first_line_chars':2,'status':'authored and OOXML checked; rendering still required'}
    (ROOT/'reports/stage5_05_local/document_build.json').write_text(json.dumps(result,indent=2)+'\n')
    print('DOCX_BUILT',json.dumps(result))
if __name__=='__main__':main()
