"""Audit RC manual bytes, source excerpts, images, styles and actual visual ledger."""
from pathlib import Path
import hashlib,json,re,zipfile,io,shutil
from collections import Counter
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
import stage5_05_manual as excerpts

ROOT=Path(__file__).resolve().parents[1]
RC=ROOT/'releases/stage5-rc1'
REPORT=ROOT/'reports/stage5_rc1_local'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
 md=RC/'docs/RK3576_STAGE5_ENGINEERING_MANUAL.md';word=md.with_suffix('.docx')
 old=ROOT/'docs'/word.name;doc=Document(word);original=Document(old)
 assert sha(md)==sha(ROOT/'docs'/md.name)
 expected=[];fence=False
 for line in md.read_text(encoding='utf-8').splitlines():
  if line.startswith('```'):fence=not fence
  elif fence:expected.append(line)
 assert not fence
 actual=[p.text for p in doc.paragraphs if p.style.name=='Source Code']
 assert expected==actual and len(actual)==4449
 def content(d):return [p.text for p in d.paragraphs if not p.style.name.lower().startswith('toc')]
 assert content(doc)==content(original),'non-TOC content changed'
 assert [[c.text for row in t.rows for c in row.cells] for t in doc.tables]==[[c.text for row in t.rows for c in row.cells] for t in original.tables]
 assert all(abs(s.page_width.mm-210)<0.05 and abs(s.page_height.mm-297)<0.05 for s in doc.sections)
 normal=doc.styles['Normal'];family=normal.element.get_or_add_rPr().find(qn('w:rFonts'))
 assert family.get(qn('w:ascii'))=='Times New Roman' and family.get(qn('w:eastAsia')) in {'楷体','KaiTi'}
 assert normal.font.size.pt==12 and normal.paragraph_format.line_spacing==1.5
 assert normal.element.get_or_add_pPr().find(qn('w:ind')).get(qn('w:firstLineChars'))=='200'
 def inherited(s,getter):
  while s is not None:
   value=getter(s)
   if value is not None:return value
   s=s.base_style
 for name in ['Title',*[f'Heading {i}' for i in range(1,7)]]:
  s=doc.styles[name];assert inherited(s,lambda x:x.font.bold) and inherited(s,lambda x:x.font.size).pt==12
  color=s.element.get_or_add_rPr().find(qn('w:color'));assert color is None or color.get(qn('w:val'))=='000000'
 manifest=json.loads((ROOT/'reports/stage5_05_local/manual_source_manifest.json').read_text(encoding='utf-8'))
 for r in manifest:
  assert sha(RC/'source'/r['file'])==r['file_sha256']==sha(ROOT/r['file'])
  assert hashlib.sha256(excerpts.block(r['file'],r['symbol']).encode()).hexdigest()==r['extracted_sha256']
 def media(path):
  with zipfile.ZipFile(path) as z:
   result={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
   for n in result:Image.open(io.BytesIO(z.read(n))).verify()
   assert z.testzip() is None
   return result
 assert media(word)==media(old) and len(media(word))==11
 assert len(doc.inline_shapes)==len(original.inline_shapes)==14
 render=ROOT/'reports/stage5_05_local/render_rc1_a4';info=json.loads((render/'result.json').read_text())
 ledger=json.loads((REPORT/'visual_review.json').read_text());assert ledger['inspected_pages']==list(range(1,239)) and ledger['clean']
 assert info['pages']==238 and info['docx_sha256']==sha(word)
 info['visual_review_complete']=True
 page_hashes={str(i):sha(render/'pages'/f'page-{i}.png') for i in ledger['inspected_pages']}
 ledger['page_sha256']=page_hashes;(REPORT/'visual_review.json').write_text(json.dumps(ledger,indent=2)+'\n')
 interview={}
 for stage in [3,4,5]:
  p=ROOT/f'docs/stage5_0{stage}_interview.md'
  interview[f'5.{stage}']=len(re.findall(r'^##\s+(?:操作\s+)?\d+',p.read_text(encoding='utf-8'),re.M))
 result={'version':json.loads((RC/'VERSION.json').read_text())['source_snapshot_id'],'pages':238,'paper':'A4 210 x 297 mm','all_pages_inspected':True,'clean':True,'docx_sha256':sha(word),'markdown_sha256':sha(md),'historical_docx_sha256':sha(old),'non_TOC_paragraphs_and_tables_unchanged':True,'source_code_lines':4449,'source_excerpts':len(manifest),'all_excerpts_match_current_RC_source':True,'image_placements':14,'unique_embedded_images':11,'embedded_images_unchanged_and_decodable':True,'body_Chinese':'楷体 12pt','body_Latin':'Times New Roman 12pt','line_spacing':1.5,'first_line_chars':2,'headings_bold_black':True,'interview_later_sections':interview,'interview_total_inherited':61,'render':info,'limitations':'Document and source checks do not establish new hardware outcomes. Original 252-page Letter manual retained in historical evidence; RC A4 reflow is 238 pages.'}
 assert interview=={'5.3':15,'5.4':14,'5.5':12}
 (REPORT/'document_QA.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
 for name in ['document_QA.json','visual_review.json']:shutil.copyfile(REPORT/name,RC/'reports'/name)
 for p in render.glob('*'):
  if p.is_file() and p.suffix in {'.json','.log'}:
   target=RC/'reports/manual_render'/p.name;target.parent.mkdir(exist_ok=True);shutil.copyfile(p,target)
 print(json.dumps(result,ensure_ascii=False))
if __name__=='__main__':main()
